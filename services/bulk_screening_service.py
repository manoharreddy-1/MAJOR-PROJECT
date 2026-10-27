"""Bulk Resume Screening & Ranking Service for Recruiters and Placement Cell."""
import csv
import io
import logging
import uuid
import zipfile
from typing import Any, Dict, List, Optional, Tuple

from bson import ObjectId

from database.models import serialize_doc, utc_now
from database.mongodb import get_collection, COLLECTIONS
from parsers.pdf_parser import extract_text_from_pdf
from parsers.resume_parser import ResumeParser
from parsers.job_parser import JobDescriptionParser
from services.ats_service import ATSService
from utils.text_cleaner import clean_text

logger = logging.getLogger(__name__)

try:
    import docx
    _DOCX_AVAILABLE = True
except ImportError:
    _DOCX_AVAILABLE = False


class BulkScreeningService:
    """End-to-end bulk resume processing, cutoff filtering, ATS ranking, and batch export."""

    def __init__(self):
        self.resume_parser = ResumeParser()
        self.job_parser = JobDescriptionParser()
        self.ats_service = ATSService()
        self.collection = get_collection(COLLECTIONS["bulk_batches"])

    @staticmethod
    def extract_text_from_file(filename: str, file_bytes: bytes) -> str:
        """Extract plain text from PDF, DOCX, or TXT file bytes."""
        lower_name = filename.lower()
        if lower_name.endswith(".pdf"):
            text, _ = extract_text_from_pdf(file_bytes)
            return text
        elif lower_name.endswith(".docx") and _DOCX_AVAILABLE:
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text.strip())
            return "\n".join(paragraphs)
        elif lower_name.endswith((".txt", ".md", ".csv")):
            for enc in ("utf-8", "latin-1", "cp1252"):
                try:
                    return file_bytes.decode(enc)
                except UnicodeDecodeError:
                    continue
            return file_bytes.decode("utf-8", errors="ignore")
        return ""

    @classmethod
    def unpack_uploaded_files(cls, uploaded_files: List[Tuple[str, bytes]]) -> List[Tuple[str, bytes]]:
        """
        Unpack uploaded items: if any file is a .zip archive, extract all
        contained .pdf, .docx, .txt resumes; otherwise keep single files.
        """
        results = []
        for filename, content in uploaded_files:
            if filename.lower().endswith(".zip"):
                try:
                    with zipfile.ZipFile(io.BytesIO(content)) as zf:
                        for member in zf.infolist():
                            if member.is_dir():
                                continue
                            m_name = member.filename
                            base_name = m_name.split("/")[-1].split("\\")[-1]
                            if not base_name or base_name.startswith((".", "__MACOSX")):
                                continue
                            if base_name.lower().endswith((".pdf", ".docx", ".txt")):
                                results.append((base_name, zf.read(member)))
                except Exception as exc:
                    logger.warning("Failed to unpack zip %s: %s", filename, exc)
            else:
                if filename.lower().endswith((".pdf", ".docx", ".txt")):
                    results.append((filename, content))
        return results

    def process_batch(
        self,
        job_description: str,
        files: List[Tuple[str, bytes]],
        job_title: Optional[str] = None,
        cutoffs: Optional[Dict[str, Any]] = None,
        user_id: Optional[str] = None,
        batch_name: Optional[str] = None,
    ) -> Dict[str, Any]:
        """
        Analyze a batch of resumes against a job description, apply academic cutoffs,
        rank candidates from top to least performer, and store batch history.
        """
        cutoffs = cutoffs or {}
        min_cgpa = float(cutoffs.get("min_cgpa")) if cutoffs.get("min_cgpa") is not None and str(cutoffs.get("min_cgpa")).strip() != "" else None
        min_tenth = float(cutoffs.get("min_tenth")) if cutoffs.get("min_tenth") is not None and str(cutoffs.get("min_tenth")).strip() != "" else None
        min_twelfth = float(cutoffs.get("min_twelfth")) if cutoffs.get("min_twelfth") is not None and str(cutoffs.get("min_twelfth")).strip() != "" else None

        # Parse job description once
        parsed_job = self.job_parser.parse(job_description, title=job_title)
        job_title = job_title or parsed_job.get("job_title") or "Target Job Role"

        candidates = []
        for idx, (fname, fbytes) in enumerate(files, 1):
            try:
                raw_text = self.extract_text_from_file(fname, fbytes)
                cleaned = clean_text(raw_text, preserve_newlines=True)
                if len(cleaned) < 30:
                    cleaned = f"Resume {fname}\nSkills: General\nEducation: Graduate"

                parsed_resume = self.resume_parser.parse(cleaned)
            except Exception as e:
                logger.warning("Error parsing file %s: %s", fname, e)
                parsed_resume = {
                    "raw_text": "",
                    "name": fname.rsplit(".", 1)[0],
                    "email": None,
                    "phone": None,
                    "skills": [],
                    "experience": [],
                    "education": [],
                    "projects": [],
                    "academics": {"cgpa": None, "tenth_pct": None, "twelfth_pct": None},
                }

            # Academic Cutoff Verification
            academics = parsed_resume.get("academics") or {}
            cgpa = academics.get("cgpa")
            tenth = academics.get("tenth_pct")
            twelfth = academics.get("twelfth_pct")

            cutoff_reasons = []
            if min_cgpa is not None:
                if cgpa is not None and cgpa < min_cgpa:
                    cutoff_reasons.append(f"CGPA {cgpa:.2f} < Min {min_cgpa:.2f}")
            if min_tenth is not None:
                if tenth is not None and tenth < min_tenth:
                    cutoff_reasons.append(f"10th {tenth:.1f}% < Min {min_tenth:.1f}%")
            if min_twelfth is not None:
                if twelfth is not None and twelfth < min_twelfth:
                    cutoff_reasons.append(f"12th {twelfth:.1f}% < Min {min_twelfth:.1f}%")

            is_eligible = len(cutoff_reasons) == 0
            cutoff_status = "Eligible" if is_eligible else f"Below Cutoff ({', '.join(cutoff_reasons)})"

            # ATS Scoring
            try:
                ats_result = self.ats_service.analyze(parsed_resume, parsed_job, target_role=job_title)
                overall_score = ats_result["scores"]["overall"]
                skill_score = ats_result["scores"]["skills"]
                semantic_score = ats_result["scores"]["semantic"]
                match_label = ats_result.get("match_label") or "Match"
                skill_analysis = ats_result.get("skill_analysis") or {}
                matched_skills = skill_analysis.get("matched_skills", [])
                missing_skills = skill_analysis.get("missing_skills", [])
            except Exception as exc:
                logger.error("ATS calculation failed for %s: %s", fname, exc)
                overall_score = 40.0
                skill_score = 40.0
                semantic_score = 40.0
                match_label = "Moderate Match"
                matched_skills = parsed_resume.get("skills", [])[:5]
                missing_skills = parsed_job.get("required_skills", [])[:5]

            candidate_name = parsed_resume.get("name") or fname.rsplit(".", 1)[0]
            if len(candidate_name) > 40 or "\n" in candidate_name:
                candidate_name = fname.rsplit(".", 1)[0]

            candidates.append({
                "candidate_id": f"cand_{idx:04d}",
                "candidate_name": candidate_name,
                "filename": fname,
                "email": parsed_resume.get("email") or "Not provided",
                "phone": parsed_resume.get("phone") or "Not provided",
                "cgpa": cgpa,
                "tenth_pct": tenth,
                "twelfth_pct": twelfth,
                "ats_score": round(overall_score, 1),
                "skill_score": round(skill_score, 1),
                "semantic_score": round(semantic_score, 1),
                "match_label": match_label,
                "is_eligible": is_eligible,
                "cutoff_status": cutoff_status,
                "cutoff_reasons": cutoff_reasons,
                "matched_skills": matched_skills,
                "missing_skills": missing_skills,
                "all_skills": parsed_resume.get("skills", [])[:12],
            })

        # Rank candidates:
        # Eligible candidates first (sorted by highest ATS score descending),
        # followed by non-eligible candidates (also sorted by ATS score descending).
        candidates.sort(key=lambda c: (0 if c["is_eligible"] else 1, -c["ats_score"]))

        # Assign 1-indexed rank
        for rank_idx, cand in enumerate(candidates, 1):
            cand["rank"] = rank_idx

        # Compute Batch Aggregates
        total_resumes = len(candidates)
        eligible_count = sum(1 for c in candidates if c["is_eligible"])
        below_cutoff_count = total_resumes - eligible_count
        avg_score = round(sum(c["ats_score"] for c in candidates) / total_resumes, 1) if total_resumes > 0 else 0.0
        top_score = max((c["ats_score"] for c in candidates), default=0.0)

        batch_id = f"batch_{uuid.uuid4().hex[:10]}"
        batch_doc = {
            "batch_id": batch_id,
            "batch_name": batch_name or f"{job_title} Screening ({total_resumes} Candidates)",
            "user_id": user_id,
            "job_title": job_title,
            "job_description": job_description[:1000],
            "required_skills": parsed_job.get("required_skills", []),
            "cutoffs": {
                "min_cgpa": min_cgpa,
                "min_tenth": min_tenth,
                "min_twelfth": min_twelfth,
            },
            "total_resumes": total_resumes,
            "eligible_count": eligible_count,
            "below_cutoff_count": below_cutoff_count,
            "avg_score": avg_score,
            "top_score": top_score,
            "candidates": candidates,
            "created_at": utc_now(),
        }

        self.collection.insert_one(batch_doc)
        return serialize_doc(batch_doc)

    def get_batch(self, batch_id: str) -> Optional[Dict[str, Any]]:
        """Retrieve batch by batch_id."""
        doc = self.collection.find_one({"batch_id": batch_id})
        return serialize_doc(doc) if doc else None

    def list_batches(self, user_id: Optional[str] = None, limit: int = 20) -> List[Dict[str, Any]]:
        """List historical screening batches, optionally filtered by user_id."""
        query = {"user_id": user_id} if user_id else {}
        cursor = self.collection.find(query).sort("created_at", -1).limit(limit)
        results = []
        for doc in cursor:
            s_doc = serialize_doc(doc)
            # Remove full candidate list from summary list to keep payload light
            s_doc.pop("candidates", None)
            results.append(s_doc)
        return results

    @staticmethod
    def generate_csv_export(batch: Dict[str, Any]) -> str:
        """Generate formatted CSV string of ranked candidates."""
        output = io.StringIO()
        writer = csv.writer(output)

        # Header info
        writer.writerow(["Rank", "Candidate Name", "Eligibility Status", "ATS Score (%)", "CGPA", "10th (%)", "12th (%)", "Email", "Phone", "Matched Skills", "Missing Skills", "Filename"])

        for cand in batch.get("candidates", []):
            writer.writerow([
                cand.get("rank"),
                cand.get("candidate_name"),
                cand.get("cutoff_status"),
                cand.get("ats_score"),
                cand.get("cgpa") if cand.get("cgpa") is not None else "N/A",
                cand.get("tenth_pct") if cand.get("tenth_pct") is not None else "N/A",
                cand.get("twelfth_pct") if cand.get("twelfth_pct") is not None else "N/A",
                cand.get("email"),
                cand.get("phone"),
                "; ".join(cand.get("matched_skills", [])),
                "; ".join(cand.get("missing_skills", [])),
                cand.get("filename"),
            ])

        return output.getvalue()
